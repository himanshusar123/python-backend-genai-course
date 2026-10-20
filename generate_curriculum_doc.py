import re
import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def create_curriculum_document(md_path, docx_path):
    doc = docx.Document()
    
    # Page setup - 0.75 in margins
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)
        
        # Header / Footer
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("Shristi Tech Academy — Python Backend + GenAI Development")
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = RGBColor(148, 163, 184)
        
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("Confidential & Proprietary — Shristi Tech Academy Flagship Curriculum")
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = RGBColor(148, 163, 184)

    # Styles
    styles = doc.styles
    normal_style = styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(10.5)
    normal_style.font.color.rgb = RGBColor(30, 41, 59)
    normal_style.paragraph_format.line_spacing = 1.2
    normal_style.paragraph_format.space_after = Pt(4)

    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    in_table = False
    table_lines = []

    def flush_table():
        nonlocal in_table, table_lines
        if not table_lines:
            in_table = False
            return
        
        # Parse table lines
        parsed_rows = []
        for tline in table_lines:
            tline = tline.strip()
            if not tline or tline.startswith('| :---') or tline.startswith('|:---') or tline.startswith('|---'):
                continue
            cells = [c.strip() for c in tline.split('|')[1:-1]]
            if cells:
                parsed_rows.append(cells)
        
        if parsed_rows:
            num_cols = max(len(r) for r in parsed_rows)
            # pad shorter rows
            for r in parsed_rows:
                while len(r) < num_cols:
                    r.append("")
            
            table = doc.add_table(rows=len(parsed_rows), cols=num_cols)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = True
            
            for row_idx, r in enumerate(parsed_rows):
                is_header = (row_idx == 0)
                row = table.rows[row_idx]
                for col_idx, text in enumerate(r):
                    cell = row.cells[col_idx]
                    set_cell_margins(cell, top=120, bottom=120, left=160, right=160)
                    if is_header:
                        set_cell_background(cell, "1E293B")  # Dark slate
                    elif row_idx % 2 == 1:
                        set_cell_background(cell, "F8FAFC")  # Light gray
                    else:
                        set_cell_background(cell, "FFFFFF")
                    
                    cp = cell.paragraphs[0]
                    cp.paragraph_format.space_after = Pt(2)
                    cp.paragraph_format.line_spacing = 1.15
                    
                    # Clean markdown bold markers
                    clean_text = text.replace('**', '').replace('`', '')
                    crun = cp.add_run(clean_text)
                    if is_header:
                        crun.font.bold = True
                        crun.font.color.rgb = RGBColor(255, 255, 255)
                        crun.font.size = Pt(10)
                    else:
                        crun.font.size = Pt(9.5)
                        if '**' in text or col_idx == 0:
                            crun.font.bold = True
            
            # Add subtle spacing after table
            sp = doc.add_paragraph()
            sp.paragraph_format.space_before = Pt(4)
            sp.paragraph_format.space_after = Pt(6)

        table_lines = []
        in_table = False

    i = 0
    while i < len(lines):
        line = lines[i].rstrip('\r\n')
        stripped = line.strip()

        # Check table
        if stripped.startswith('|'):
            in_table = True
            table_lines.append(stripped)
            i += 1
            continue
        elif in_table:
            flush_table()

        if not stripped:
            i += 1
            continue

        # Check Headings
        if stripped.startswith('# '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(18)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(stripped[2:])
            run.font.size = Pt(20)
            run.font.bold = True
            run.font.color.rgb = RGBColor(30, 58, 138) # Deep navy blue
        elif stripped.startswith('## '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(stripped[3:])
            run.font.size = Pt(15)
            run.font.bold = True
            run.font.color.rgb = RGBColor(79, 70, 229) # Indigo
        elif stripped.startswith('### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(stripped[4:])
            run.font.size = Pt(12.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor(14, 116, 144) # Cyan/Dark Teal
        elif stripped.startswith('#### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(stripped[5:])
            run.font.size = Pt(11)
            run.font.bold = True
            run.font.color.rgb = RGBColor(51, 65, 85) # Slate
        elif stripped.startswith('---'):
            # Horizontal rule or divider
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            r = p.add_run("_________________________________________________________________________________")
            r.font.size = Pt(8)
            r.font.color.rgb = RGBColor(226, 232, 240)
        elif stripped.startswith('- ') or stripped.startswith('* '):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            bullet_text = stripped[2:]
            
            # parse bold segments
            parts = re.split(r'(\*\*.*?\*\*)', bullet_text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part[2:-2])
                    run.bold = True
                else:
                    # remove backticks
                    run = p.add_run(part.replace('`', ''))
        elif re.match(r'^\d+\.\s', stripped):
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            num_text = re.sub(r'^\d+\.\s*', '', stripped)
            parts = re.split(r'(\*\*.*?\*\*)', num_text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part[2:-2])
                    run.bold = True
                else:
                    run = p.add_run(part.replace('`', ''))
        elif stripped.startswith('```'):
            # code block
            code_lines = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith('```'):
                code_lines.append(lines[i].rstrip('\r\n'))
                i += 1
            code_text = "\n".join(code_lines)
            
            # Add table container for code block
            tbl = doc.add_table(rows=1, cols=1)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            c = tbl.rows[0].cells[0]
            set_cell_background(c, "F1F5F9")
            set_cell_margins(c, top=100, bottom=100, left=150, right=150)
            cp = c.paragraphs[0]
            cp.paragraph_format.space_after = Pt(0)
            crun = cp.add_run(code_text)
            crun.font.name = 'Consolas'
            crun.font.size = Pt(9)
            crun.font.color.rgb = RGBColor(15, 23, 42)
            
            sp = doc.add_paragraph()
            sp.paragraph_format.space_after = Pt(4)
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(4)
            parts = re.split(r'(\*\*.*?\*\*)', stripped)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part[2:-2])
                    run.bold = True
                else:
                    run = p.add_run(part.replace('`', ''))

        i += 1

    if in_table:
        flush_table()

    doc.save(docx_path)
    print(f"Successfully generated docx: {docx_path}")

if __name__ == "__main__":
    src_md = r"C:\Users\Himanshu Sardana\.gemini\antigravity-ide\scratch\python-genai-course-landing-page\PYTHON_BACKEND_GENAI_CURRICULUM_AND_WEBPAGE_CONTENT.md"
    out_docx = r"C:\Users\Himanshu Sardana\.gemini\antigravity-ide\scratch\python-genai-course-landing-page\Python_Backend_GenAI_Course_Curriculum_Document.docx"
    create_curriculum_document(src_md, out_docx)
